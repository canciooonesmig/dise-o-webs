"""Utilidades de construcción: carga de ficheros R, crosswalk de identificadores y escritura del libro Excel."""
import json, os, re, warnings
import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")
BUILD_DATE = "2026-10-08"  # fecha de construcción / consulta de las fuentes dinámicas


def load_rda(path):
    """Lee un .rda/.RData devolviendo el primer data.frame. Intenta pyreadr y, si falla por codificación, rdata (latin1)."""
    try:
        import pyreadr
        r = pyreadr.read_r(path)
        return next(iter(r.values()))
    except Exception:
        import rdata
        parsed = rdata.parser.parse_file(path, extension="rda")
        conv = rdata.conversion.convert(parsed, default_encoding="latin1")
        df = next(iter(conv.values()))
        df.columns = [str(c) for c in df.columns]
        return pd.DataFrame(df)


def r_date(x):
    """Convierte fechas R (días desde 1970) o datetime a texto ISO."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return ""
    if isinstance(x, (pd.Timestamp,)) or hasattr(x, "strftime"):
        try:
            return x.strftime("%Y-%m-%d")
        except Exception:
            return ""
    try:
        return (pd.Timestamp("1970-01-01") + pd.Timedelta(days=float(x))).strftime("%Y-%m-%d")
    except Exception:
        return str(x)


# ---------------------------------------------------------------------------
# Crosswalk de identificadores
# ---------------------------------------------------------------------------
COW_MICRO = {"BHM": "BHS", "DMA": "DMA", "GRN": "GRD", "SLU": "LCA", "SVG": "VCT", "AAB": "ATG", "SKN": "KNA",
             "BLZ": "BLZ", "MNC": "MCO", "LIE": "LIE", "AND": "AND", "SNM": "SMR", "BRU": "BRN", "KIR": "KIR",
             "TUV": "TUV", "TON": "TON", "NAU": "NRU", "MSI": "MHL", "PAL": "PLW", "FSM": "FSM", "WSM": "WSM"}
GW_EXTRA = {31: "BHS", 80: "BLZ", 835: "BRN", 340: "SRB", 89: "H_UPC", 99: "H_GCL", 563: "H_TRA", 564: "H_OFS",
            711: "TBT", 815: "H_ANNAM"}
EXTRA_ENTITIES = [
    # id, nombre_es, tipo, inicio, fin, nota
    ("H_UPC", "Provincias Unidas del Centro de América", "estado_historico", "1823-07-01", "1839-12-31", "Gleditsch-Ward 89"),
    ("H_GCL", "Gran Colombia", "estado_historico", "1821-08-30", "1830-09-22", "Gleditsch-Ward 99"),
    ("H_TRA", "Transvaal (República Sudafricana)", "estado_historico", "1852-01-01", "1910-05-30", "Gleditsch-Ward 563"),
    ("H_OFS", "Estado Libre de Orange", "estado_historico", "1854-03-28", "1910-05-30", "Gleditsch-Ward 564"),
    ("TBT", "Tíbet", "estado_historico", "1913-01-01", "1950-10-01", "Gleditsch-Ward 711; incorporado a la RPC"),
    ("H_ANNAM", "Vietnam (Annam/Cochinchina/Tonkín)", "estado_historico", "1816-01-01", "1893-01-01", "Gleditsch-Ward 815"),
    ("SKM", "Sikkim", "estado_historico", "1947", "1975-05-16", "Protectorado indio; anexionado en 1975"),
    ("CHI", "Islas del Canal (Jersey y Guernsey, agregado WDI)", "territorio_dependiente_o_especial", "", "", "Código WDI CHI; ver JEY y GGY"),
    ("WLD", "Mundo (agregado)", "agregado_global", "", "", "Pseudoentidad para acontecimientos globales"),
    ("EUU", "Unión Europea", "organizacion_supranacional", "1958-01-01", "", "Código WDI EUU"),
]
VDEM_REGION6 = {1: "Europa del Este y Asia Central", 2: "América Latina y el Caribe", 3: "Oriente Medio y Norte de África",
                4: "África Subsahariana", 5: "Europa Occidental y Norteamérica", 6: "Asia y Pacífico"}
ISO_FIX = {"UNK": "XKX"}
ES_NAMES = {"DDR": "República Democrática Alemana", "VDR": "República de Vietnam (Vietnam del Sur)",
            "YMD": "Yemen del Sur (RDP de Yemen)", "SML": "Somalilandia", "ZZB": "Zanzíbar", "PSG": "Palestina/Gaza",
            "PSB": "Palestina/Mandato británico", "BVR": "Baviera", "SAX": "Sajonia", "WRG": "Wurtemberg",
            "HVR": "Hannover", "BDN": "Baden", "PPS": "Estados Pontificios", "TWS": "Reino de las Dos Sicilias",
            "SPD": "Piamonte-Cerdeña", "TSC": "Toscana", "MDN": "Módena", "PRM": "Parma", "HKS": "Hesse-Kassel",
            "HDM": "Hesse-Darmstadt", "MCL": "Mecklemburgo-Schwerin", "NSS": "Nassau", "OLD": "Oldemburgo",
            "BRW": "Brunswick", "SXN": "Sajonia-Weimar-Eisenach", "HRG": "Hamburgo"}
WDI_GEO_FIX = {"KOS": "XKX", "DEU_EAST": "DDR", "YEM_SOUTH": "YMD", "GBM": "IMN", "NLD_CURACAO": "CUW", "CHANISL": "CHI"}


def build_crosswalk(raw):
    """Devuelve (master DataFrame, ext_ids DataFrame, cow->id dict, gw->id dict)."""
    vd = pd.read_pickle(os.path.join(raw, "vdem_subset.pkl"))
    vm = vd.groupby("country_text_id").agg(vdem_name=("country_name", "first"), vdem_num=("country_id", "first"),
                                          cow=("COWcode", "first"), y0=("year", "min"), y1=("year", "max"),
                                          reg6=("e_regionpol_6C", "last")).reset_index()
    cow_to_id = {int(r.cow): r.country_text_id for r in vm.itertuples() if pd.notna(r.cow)}
    cow = load_rda(os.path.join(raw, "ps/cow_states.rda"))
    for r in cow.itertuples():
        if int(r.ccode) not in cow_to_id and r.stateabb in COW_MICRO:
            cow_to_id[int(r.ccode)] = COW_MICRO[r.stateabb]
    cow_to_id.update({260: "DEU", 305: "AUT", 316: "CZE", 679: "YEM", 732: "KOR"})
    gw = load_rda(os.path.join(raw, "ps/gw_states.rda"))
    gw_to_id = {}
    for r in gw.itertuples():
        c = int(r.gwcode)
        if c in GW_EXTRA:
            gw_to_id[c] = GW_EXTRA[c]
        elif c in cow_to_id:
            gw_to_id[c] = cow_to_id[c]
    # Correcciones GW conocidas (los códigos GW difieren de COW en algunos casos)
    gw_to_id.update({260: "DEU", 265: "DDR", 678: "YEM", 680: "YMD", 345: "SRB", 347: "XKX", 816: "VNM", 817: "VDR"})

    cj = json.load(open(os.path.join(raw, "countries.json")))
    rows = {}
    for c in cj:
        cid = ISO_FIX.get(c["cca3"], c["cca3"])
        spa = c.get("translations", {}).get("spa", {})
        indep = c.get("independent")
        un = c.get("unMember")
        if cid in ("TWN", "XKX", "PSE"):
            etype = "estado_reconocimiento_limitado"
        elif cid == "ESH":
            etype = "territorio_no_autonomo_disputado"
        elif indep and un:
            etype = "estado_soberano_miembro_ONU"
        elif indep:
            etype = "estado_soberano_no_miembro_ONU"
        else:
            etype = "territorio_dependiente_o_especial"
        rows[cid] = dict(country_id=cid, name_es=spa.get("common", c["name"]["common"]), name_official_es=spa.get("official", ""),
                         name_en=c["name"]["common"], name_official_en=c["name"]["official"], entity_type=etype,
                         iso_alpha2=c.get("cca2", ""), iso_alpha3=c["cca3"] if c["cca3"] != "UNK" else "",
                         iso_numeric=c.get("ccn3", ""), un_member=bool(un), independent_flag=indep,
                         capital=";".join(c.get("capital", []) or []), region=c.get("region", ""), subregion=c.get("subregion", ""),
                         area_km2=c.get("area"), landlocked=c.get("landlocked"),
                         borders=";".join(ISO_FIX.get(b, b) for b in c.get("borders", [])),
                         n_land_borders=len(c.get("borders", [])),
                         lat=(c.get("latlng") or [None, None])[0], lon=(c.get("latlng") or [None, None])[1],
                         currencies=";".join(c.get("currencies", {}).keys()) if isinstance(c.get("currencies"), dict) else "",
                         languages=";".join(c.get("languages", {}).values()) if isinstance(c.get("languages"), dict) else "",
                         geo_source="SRC_MLEDOZE")
    for r in vm.itertuples():
        if r.country_text_id not in rows:
            hist = r.y1 < 2025
            rows[r.country_text_id] = dict(country_id=r.country_text_id, name_es=r.vdem_name, name_en=r.vdem_name,
                                           entity_type="estado_historico" if hist else "entidad_autonoma_o_no_reconocida",
                                           region=VDEM_REGION6.get(int(r.reg6) if pd.notna(r.reg6) else -1, ""),
                                           geo_source="SRC_VDEM")
    for eid, name, et, s, e, note in EXTRA_ENTITIES:
        if eid not in rows:
            rows[eid] = dict(country_id=eid, name_es=name, name_en=name, entity_type=et, entity_note=note,
                             sovereignty_start_curated=s, sovereignty_end_curated=e, geo_source="CURADO")
    for k, v in ES_NAMES.items():
        if k in rows:
            rows[k]["name_es"] = v
    master = pd.DataFrame(rows.values())
    # Atributos V-Dem
    vmi = vm.set_index("country_text_id")
    master["vdem_id"] = master.country_id.map(lambda x: x if x in vmi.index else "")
    master["vdem_country_num"] = master.country_id.map(lambda x: vmi.vdem_num.get(x))
    master["vdem_first_year"] = master.country_id.map(lambda x: vmi.y0.get(x))
    master["vdem_last_year"] = master.country_id.map(lambda x: vmi.y1.get(x))
    master["region_vdem6"] = master.country_id.map(lambda x: VDEM_REGION6.get(int(vmi.reg6.get(x)), "") if x in vmi.index and pd.notna(vmi.reg6.get(x)) else "")

    # Identificadores externos con vigencia
    ext = []
    for r in cow.itertuples():
        cid = cow_to_id.get(int(r.ccode), "")
        ext.append(dict(country_id=cid, scheme="COW_ccode", code=str(int(r.ccode)), code_label=r.statenme,
                        abbrev=r.stateabb, valid_from=f"{int(r.styear):04d}-{int(r.stmonth):02d}-{int(r.stday):02d}",
                        valid_to=f"{int(r.endyear):04d}-{int(r.endmonth):02d}-{int(r.endday):02d}",
                        source_id="SRC_COW_STATES", note="COW 2016: fin=2016-12-31 indica que seguía vigente al cierre de la serie"))
    for r in gw.itertuples():
        cid = gw_to_id.get(int(r.gwcode), "")
        ext.append(dict(country_id=cid, scheme="GW_code", code=str(int(r.gwcode)), code_label=r.statename, abbrev=r.stateabb,
                        valid_from=r_date(r.startdate), valid_to=r_date(r.enddate), source_id="SRC_GW_STATES",
                        note="Gleditsch-Ward: fin=2017-12-31 indica vigencia al cierre de la serie"))
    for r in vm.itertuples():
        ext.append(dict(country_id=r.country_text_id, scheme="VDEM_country_text_id", code=r.country_text_id,
                        code_label=r.vdem_name, abbrev=str(int(r.vdem_num)) if pd.notna(r.vdem_num) else "",
                        valid_from=str(int(r.y0)), valid_to=str(int(r.y1)), source_id="SRC_VDEM", note=""))
    for c in rows.values():
        if c.get("iso_alpha3"):
            ext.append(dict(country_id=c["country_id"], scheme="ISO3166_alpha3", code=c["iso_alpha3"], code_label=c["name_en"],
                            abbrev=c.get("iso_alpha2", ""), valid_from="", valid_to="", source_id="SRC_MLEDOZE", note=""))
    ext = pd.DataFrame(ext)
    # Inicio/fin de soberanía según COW y GW
    cw = ext[ext.scheme == "COW_ccode"].groupby("country_id").agg(cow_first=("valid_from", "min"), cow_last=("valid_to", "max"))
    gwx = ext[ext.scheme == "GW_code"].groupby("country_id").agg(gw_first=("valid_from", "min"), gw_last=("valid_to", "max"))
    master = master.merge(cw, left_on="country_id", right_index=True, how="left").merge(gwx, left_on="country_id", right_index=True, how="left")
    master["cow_codes"] = master.country_id.map(ext[ext.scheme == "COW_ccode"].groupby("country_id").code.apply(lambda s: ";".join(sorted(set(s), key=int))))
    master["gw_codes"] = master.country_id.map(ext[ext.scheme == "GW_code"].groupby("country_id").code.apply(lambda s: ";".join(sorted(set(s), key=int))))
    return master, ext, cow_to_id, gw_to_id


# ---------------------------------------------------------------------------
# Escritura del libro
# ---------------------------------------------------------------------------
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LOW_FILL = PatternFill("solid", fgColor="F8CBAD")
MID_FILL = PatternFill("solid", fgColor="FFF2CC")
MISS_FILL = PatternFill("solid", fgColor="D9D9D9")
CONF_COLS = {"confidence_level", "certainty", "confidence"}
URL_COLS = {"url", "official_url", "source_url"}
LONG_TEXT_HINT = ("description", "notes", "note", "nota", "consequences", "antecedents", "causes", "mechanism",
                  "evidence", "provisions", "methodology", "limitations", "situation", "conditions", "debate",
                  "objectives", "outcome", "definition", "purpose", "question", "clarification", "responses", "text",
                  "cases", "threats", "opportunities", "relations", "events", "assumptions", "role", "products")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
        return None
    if isinstance(v, (np.floating,)):
        v = float(v)
        return None if (np.isnan(v) or np.isinf(v)) else v
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if v == "" or v.lower() == "nan":
            return None
        if v[:1] in ("=", "+", "-", "@") and not re.match(r"^-?\d", v):
            v = "'" + v  # evitar interpretación como fórmula
        return v[:32000]
    if pd.isna(v) if not isinstance(v, (list, dict)) else False:
        return None
    return v


def write_book(path, tables):
    """tables: lista de dict con name, df, (opcional) number_formats, widths."""
    wb = Workbook(write_only=True)
    for t in tables:
        df = t["df"]
        ws = wb.create_sheet(t["name"][:31])
        cols = list(df.columns)
        ncol = len(cols)
        # anchos
        for i, c in enumerate(cols, 1):
            if c in URL_COLS:
                w = 45
            elif any(h in c.lower() for h in LONG_TEXT_HINT):
                w = 50
            else:
                sample = df[c].astype(str).head(200).str.len()
                w = int(min(max(len(c) + 2, (sample.quantile(0.9) if len(sample) else 8) + 2, 8), 40))
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2" if t.get("freeze_col") else "A2"
        # cabecera
        hdr = []
        for c in cols:
            cell = WriteOnlyCell(ws, value=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            hdr.append(cell)
        ws.append(hdr)
        # tipos por columna
        date_cols = set()
        for c in cols:
            s = df[c].dropna()
            s = s[s.astype(str) != ""]
            if len(s) and s.dtype == object and s.astype(str).str.match(ISO_DATE).all() and ("date" in c or c.endswith(("_start", "_end", "start", "end"))):
                date_cols.add(c)
        wrap_cols = {c for c in cols if any(h in c.lower() for h in LONG_TEXT_HINT)}
        numfmt = t.get("number_formats", {})
        int_cols = {c for c in cols if c in ("year", "year_start", "year_end") or c.endswith("_year")}
        records = df.itertuples(index=False, name=None)
        for rec in records:
            row = []
            for c, v in zip(cols, rec):
                v = _clean(v)
                if v is None:
                    row.append(None)
                    continue
                if c in date_cols:
                    try:
                        cell = WriteOnlyCell(ws, value=pd.Timestamp(v).to_pydatetime())
                        cell.number_format = "yyyy-mm-dd"
                        row.append(cell)
                        continue
                    except Exception:
                        pass
                if c in URL_COLS and isinstance(v, str) and v.startswith("http"):
                    cell = WriteOnlyCell(ws, value=v)
                    cell.hyperlink = v
                    cell.font = Font(color="0563C1", underline="single")
                    row.append(cell)
                    continue
                if c in wrap_cols and isinstance(v, str) and len(v) > 50:
                    cell = WriteOnlyCell(ws, value=v)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    row.append(cell)
                    continue
                if isinstance(v, float):
                    cell = WriteOnlyCell(ws, value=v)
                    cell.number_format = numfmt.get(c, "#,##0.0##" if abs(v) < 1e6 else "#,##0")
                    row.append(cell)
                    continue
                if isinstance(v, int) and not isinstance(v, bool) and c not in int_cols and abs(v) >= 10000:
                    cell = WriteOnlyCell(ws, value=v)
                    cell.number_format = "#,##0"
                    row.append(cell)
                    continue
                row.append(v)
            ws.append(row)
        nrow = len(df) + 1
        ref = f"A1:{get_column_letter(ncol)}{max(nrow, 2)}"
        tname = "T_" + re.sub(r"[^A-Za-z0-9_]", "_", t["name"])[:250]
        tab = Table(displayName=tname, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)
        # formato condicional y validación
        for i, c in enumerate(cols, 1):
            L = get_column_letter(i)
            rng = f"{L}2:{L}{max(nrow, 2)}"
            if c in CONF_COLS:
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Baja"'], fill=LOW_FILL))
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Media"'], fill=MID_FILL))
                dv = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=True)
                dv.add(rng)
                ws.data_validations.append(dv)
            if c in t.get("required_cols", []):
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEN(TRIM({L}2))=0'], fill=MISS_FILL))
            if c == "verification_status":
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CONOC_EXPERTO"'], fill=MID_FILL))
            if c in ("observation_type",) and t.get("obs_types"):
                dv = DataValidation(type="list", formula1='"' + ",".join(t["obs_types"]) + '"', allow_blank=True)
                dv.add(rng)
                ws.data_validations.append(dv)
    wb.save(path)
